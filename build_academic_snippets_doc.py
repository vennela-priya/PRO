from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path("BrainTumorAI_Academic_Code_Snippets_Report.docx").resolve()


def shade(element, fill):
    tc_pr = element._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color="000000", size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)


def add_code(doc, code):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    shade(cell, "F3F6FA")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    for line_no, line in enumerate(code.strip("\n").splitlines()):
        if line_no:
            p.add_run("\n")
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(8.2)
        run.font.color.rgb = RGBColor(32, 42, 56)
    doc.add_paragraph()


def add_snippet(doc, file_name, purpose, importance, code):
    p = doc.add_paragraph()
    p.style = "Heading 3"
    p.add_run(file_name)

    meta = doc.add_table(rows=2, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.columns[0].width = Cm(3.2)
    meta.columns[1].width = Cm(12.5)
    for row in meta.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_text(meta.cell(0, 0), "Module Purpose", bold=True, color="FFFFFF")
    shade(meta.cell(0, 0), "1F4E79")
    set_cell_text(meta.cell(0, 1), purpose)
    set_cell_text(meta.cell(1, 0), "Academic Relevance", bold=True, color="FFFFFF")
    shade(meta.cell(1, 0), "1F4E79")
    set_cell_text(meta.cell(1, 1), importance)
    doc.add_paragraph()
    add_code(doc, code)


def add_note(doc, text):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade(cell, "EAF3F8")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.font.name = "Aptos"
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(31, 78, 121)
    doc.add_paragraph()


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    for name, size, color in [
        ("Heading 1", 16, "1F4E79"),
        ("Heading 2", 13, "2F5597"),
        ("Heading 3", 11, "1F4E79"),
    ]:
        st = styles[name]
        st.font.name = "Aptos Display"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("BrainTumorAI\nAcademic Code Snippets Report")
    run.font.name = "Aptos Display"
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(31, 78, 121)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("Selected implementation snippets for final-year project documentation")
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(90, 90, 90)

    doc.add_paragraph()
    summary = doc.add_table(rows=5, cols=2)
    summary.alignment = WD_TABLE_ALIGNMENT.CENTER
    rows = [
        ("Project Type", "AI-based Brain Tumor MRI Classification and Segmentation System"),
        ("Primary Stack", "Python, PyTorch, OpenCV, FastAPI, Gradio, ReportLab"),
        ("Classes", "glioma, meningioma, notumor, pituitary"),
        ("Storage Model", "Folder-based image dataset with CSV metadata; no SQL database"),
        ("Documentation Focus", "Dataset pipeline, model architecture, inference, XAI, segmentation, API, UI, validation"),
    ]
    for i, (k, v) in enumerate(rows):
        set_cell_text(summary.cell(i, 0), k, bold=True, color="FFFFFF")
        shade(summary.cell(i, 0), "1F4E79")
        set_cell_text(summary.cell(i, 1), v)

    doc.add_page_break()

    doc.add_heading("1. System Architecture", level=1)
    add_note(doc, "The project does not contain authentication, SQL database, admin, student, or teacher workflow modules. It is an AI medical-imaging project using image folders, CSV metadata, trained checkpoints, API endpoints, and a Gradio interface.")
    add_snippet(
        doc,
        "config.yaml",
        "Central configuration for dataset, model, training, inference, and UI.",
        "Shows reproducibility and separates project settings from source code.",
        """
project:
  name: BrainTumorAI
  version: "1.0.0"
  seed: 42
  device: "cpu"

data:
  train_dir: "Project/Training"
  test_dir: "Project/Testing"
  processed_dir: "data/processed"
  num_classes: 4
  class_names: ["glioma", "meningioma", "notumor", "pituitary"]
  image_size: 224
  val_split: 0.20
  holdout_csv: "data/processed/test_holdout.csv"
  trainval_csv: "data/processed/trainval.csv"

inference:
  latency_budget_ms: 500
  tta_enabled: true
  tta_n_augments: 5
""",
    )

    doc.add_heading("2. Database Design / Dataset Storage", level=1)
    add_snippet(
        doc,
        "prepare_data.py",
        "Scans class-wise MRI folders and generates CSV metadata.",
        "Acts as the dataset indexing layer in place of a relational database.",
        """
def scan_folder(root: Path, class_names: list[str]) -> pd.DataFrame:
    records = []

    for label, cls in enumerate(class_names):
        cls_dir = root / cls
        if not cls_dir.exists():
            logger.warning(f"Class folder not found: {cls_dir}")
            continue

        for img_path in cls_dir.glob("*"):
            if img_path.suffix.lower() in {".jpg", ".jpeg", ".png", ".bmp"}:
                records.append({
                    "path": str(img_path),
                    "label": label,
                    "class_name": cls,
                    "patient_id": f"{cls}_{img_path.stem}",
                    "source": root.name,
                })

    return pd.DataFrame(records)
""",
    )
    add_snippet(
        doc,
        "prepare_data.py",
        "Creates stratified train/validation split and locked holdout test CSV.",
        "Demonstrates fair validation design and avoids mixing train/test data.",
        """
train_df, val_df = train_test_split(
    trainval_df,
    test_size=val_split,
    stratify=trainval_df["label"],
    random_state=seed,
)

train_df["split"] = "train"
val_df["split"] = "val"

trainval_tagged = pd.concat([train_df, val_df]).reset_index(drop=True)
trainval_tagged.to_csv(out_dir / "trainval.csv", index=False)
holdout_df.to_csv(out_dir / "test_holdout.csv", index=False)
""",
    )

    doc.add_heading("3. Core Functional Modules", level=1)
    add_snippet(
        doc,
        "train.py",
        "Custom PyTorch Dataset class for MRI image loading.",
        "Connects CSV metadata to model-ready tensors and supports class balancing.",
        """
class MRIDataset(Dataset):
    def __init__(self, df: pd.DataFrame, transform=None):
        self.df = df.reset_index(drop=True)
        self.transform = transform

    def __len__(self):
        return len(self.df)

    def __getitem__(self, idx):
        row = self.df.iloc[idx]
        image = cv2.imread(str(row["path"]))
        if image is None:
            image = np.zeros((224, 224, 3), dtype=np.uint8)

        image = cv2.cvtColor(image, cv2.COLOR_BGR2RGB)
        if self.transform:
            image = self.transform(image=image)["image"]

        return image, int(row["label"])

    def class_weights(self):
        counts = self.df["label"].value_counts().sort_index().values.astype(float)
        weights = 1.0 / counts
        weights /= weights.sum()
        return torch.tensor(weights, dtype=torch.float32)
""",
    )
    add_snippet(
        doc,
        "train.py",
        "MixUp and CutMix augmentation functions.",
        "Improves generalization for medical image classification.",
        """
def mixup(x, y, alpha=0.3):
    lam = np.random.beta(alpha, alpha)
    index = torch.randperm(x.size(0))
    mixed_x = lam * x + (1 - lam) * x[index]
    return mixed_x, y, y[index], lam


def cutmix(x, y, alpha=0.4):
    lam = np.random.beta(alpha, alpha)
    index = torch.randperm(x.size(0))
    _, _, h, w = x.shape
    cut_ratio = np.sqrt(1 - lam)
    cx, cy = np.random.randint(w), np.random.randint(h)

    x1 = max(0, int(cx - w * cut_ratio / 2))
    x2 = min(w, int(cx + w * cut_ratio / 2))
    y1 = max(0, int(cy - h * cut_ratio / 2))
    y2 = min(h, int(cy + h * cut_ratio / 2))

    mixed_x = x.clone()
    mixed_x[:, :, y1:y2, x1:x2] = x[index, :, y1:y2, x1:x2]
    lam = 1 - ((x2 - x1) * (y2 - y1) / (w * h))
    return mixed_x, y, y[index], lam
""",
    )

    doc.add_heading("4. AI Integration", level=1)
    add_snippet(
        doc,
        "train.py",
        "EfficientNet-based classifier.",
        "Shows transfer learning with a pretrained CNN backbone and custom head.",
        """
def make_efficientnet(num_classes, pretrained=True):
    backbone = timm.create_model(
        "efficientnet_b3",
        pretrained=pretrained,
        num_classes=0,
        global_pool="avg"
    )

    head = nn.Sequential(
        nn.Dropout(0.3),
        nn.Linear(backbone.num_features, 512),
        nn.ReLU(),
        nn.Dropout(0.2),
        nn.Linear(512, num_classes)
    )

    class EfficientNetModel(nn.Module):
        def __init__(self):
            super().__init__()
            self.backbone = backbone
            self.head = head

        def forward(self, x):
            return self.head(self.backbone(x))

    return EfficientNetModel()
""",
    )
    add_snippet(
        doc,
        "train.py",
        "Soft-voting ensemble with temperature scaling.",
        "Combines multiple classifiers and improves confidence calibration.",
        """
class Ensemble(nn.Module):
    def __init__(self, models: dict, weights: dict):
        super().__init__()
        self.models = nn.ModuleDict(models)
        self.weights = weights
        self.temperatures = nn.ParameterDict({
            name: nn.Parameter(torch.ones(1))
            for name in models
        })

    def forward(self, x):
        output = None

        for name, model in self.models.items():
            logits = model(x) / self.temperatures[name].clamp(min=0.1)
            probabilities = torch.softmax(logits, dim=-1) * self.weights[name]
            output = probabilities if output is None else output + probabilities

        return output
""",
    )

    doc.add_heading("5. Segmentation and Tumor Localization", level=1)
    add_snippet(
        doc,
        "segmentation/unet.py",
        "U-Net model for tumor mask generation.",
        "Represents encoder-decoder segmentation logic with skip connections.",
        """
class UNet(nn.Module):
    def __init__(self, in_channels=3, out_channels=1, base_filters=16):
        super().__init__()
        f = base_filters
        self.pool = nn.MaxPool2d(2)

        self.enc1 = ConvBlock(in_channels, f)
        self.enc2 = ConvBlock(f, f * 2)
        self.enc3 = ConvBlock(f * 2, f * 4)
        self.enc4 = ConvBlock(f * 4, f * 8)
        self.bottleneck = ConvBlock(f * 8, f * 16)

        self.up4 = nn.Upsample(scale_factor=2, mode="bilinear", align_corners=True)
        self.dec4 = ConvBlock(f * 16 + f * 8, f * 8)
        self.up3 = nn.Upsample(scale_factor=2, mode="bilinear", align_corners=True)
        self.dec3 = ConvBlock(f * 8 + f * 4, f * 4)
        self.up2 = nn.Upsample(scale_factor=2, mode="bilinear", align_corners=True)
        self.dec2 = ConvBlock(f * 4 + f * 2, f * 2)
        self.up1 = nn.Upsample(scale_factor=2, mode="bilinear", align_corners=True)
        self.dec1 = ConvBlock(f * 2 + f, f)
        self.outc = nn.Conv2d(f, out_channels, 1)

    def forward(self, x):
        e1 = self.enc1(x)
        e2 = self.enc2(self.pool(e1))
        e3 = self.enc3(self.pool(e2))
        e4 = self.enc4(self.pool(e3))
        b = self.bottleneck(self.pool(e4))

        d4 = self.dec4(torch.cat([e4, self.up4(b)], dim=1))
        d3 = self.dec3(torch.cat([e3, self.up3(d4)], dim=1))
        d2 = self.dec2(torch.cat([e2, self.up2(d3)], dim=1))
        d1 = self.dec1(torch.cat([e1, self.up1(d2)], dim=1))
        return self.outc(d1)
""",
    )

    doc.add_heading("6. Explainable AI", level=1)
    add_snippet(
        doc,
        "inference/gradcam.py",
        "Grad-CAM heatmap generation.",
        "Explains model prediction by highlighting influential MRI regions.",
        """
def _single_gradcam(model, target_layer, tensor, pred_class):
    hooks = _Hooks(target_layer)
    model.eval()
    model.zero_grad()

    try:
        with torch.enable_grad():
            output = model(tensor)
            probabilities = torch.softmax(output.float(), dim=1)
            score = output[0, pred_class]
            score.backward()
    finally:
        hooks.remove()

    gradients = hooks.gradients
    activations = hooks.activations
    weights = gradients.mean(dim=[2, 3], keepdim=True)

    cam = (weights * activations).sum(dim=1, keepdim=True)
    cam = F.relu(cam).squeeze().detach().cpu().numpy()

    cam_min, cam_max = float(cam.min()), float(cam.max())
    if cam_max - cam_min > 1e-8:
        cam = (cam - cam_min) / (cam_max - cam_min)
    else:
        cam = np.zeros_like(cam)

    return cam.astype(np.float32), probabilities.squeeze(0).detach().cpu().numpy()
""",
    )

    doc.add_heading("7. Backend Implementation", level=1)
    add_snippet(
        doc,
        "src/api/fastapi_app.py",
        "FastAPI endpoint for MRI prediction.",
        "Represents backend model serving and API-level image validation.",
        """
@app.post("/predict", tags=["Inference"])
async def predict(file: UploadFile = File(...),
                  include_xai: bool = True,
                  include_seg: bool = True):
    if _engine is None:
        raise HTTPException(status_code=503, detail="Model not loaded.")

    try:
        contents = await file.read()
        image = Image.open(io.BytesIO(contents)).convert("RGB")
        image_np = np.array(image)
    except Exception as e:
        raise HTTPException(status_code=422, detail=f"Cannot decode image: {e}")

    sample_id = Path(file.filename).stem if file.filename else "upload"
    result = _engine.predict(
        image=image_np,
        sample_id=sample_id,
        include_xai=include_xai,
        include_seg=include_seg,
    )

    return JSONResponse(content=result)
""",
    )
    add_snippet(
        doc,
        "src/api/inference.py",
        "End-to-end prediction pipeline.",
        "Shows preprocessing, classification, segmentation, XAI, and latency tracking.",
        """
def predict(self, image, sample_id="sample", include_xai=True, include_seg=True):
    start_time = time.perf_counter()

    tensor, resized_image = self.preprocess(image)
    preprocessing_time = time.perf_counter()

    pred_class, confidence, probabilities = self.classify(tensor, use_tta=True)
    classification_time = time.perf_counter()

    segmentation = self.segment(tensor) if include_seg else None
    segmentation_time = time.perf_counter()

    explanation = None
    if include_xai:
        explanation = self.explain(tensor, resized_image, pred_class, sample_id)

    xai_time = time.perf_counter()

    return {
        "sample_id": sample_id,
        "prediction": {
            "class_index": pred_class,
            "class_name": self.class_names[pred_class],
            "confidence": round(confidence, 4),
            "probabilities": dict(zip(self.class_names, probabilities)),
        },
        "segmentation": segmentation,
        "xai": {"gradcam_paths": explanation if explanation else {}},
        "latency_ms": {"total": round((xai_time - start_time) * 1000, 2)},
    }
""",
    )

    doc.add_heading("8. Frontend Implementation", level=1)
    add_snippet(
        doc,
        "app.py",
        "Gradio interface for MRI upload, diagnosis, Grad-CAM, segmentation, and report generation.",
        "Represents the main interactive user interface used in project demonstration.",
        """
def build_ui(model_obj, model_mode, gradcam_engine=None, segmenter=None):
    with gr.Blocks(title="NeuroScan AI", theme=gr.themes.Base()) as demo:
        result_state = gr.State(None)
        image_state = gr.State(None)
        gradcam_state = gr.State(None)
        segmentation_state = gr.State(None)

        with gr.Tabs():
            with gr.Tab("Diagnosis"):
                image_input = gr.Image(
                    label="Brain MRI Scan",
                    type="numpy",
                    height=250,
                    sources=["upload", "clipboard"],
                )
                patient_id = gr.Textbox(label="Patient ID (optional)")
                run_button = gr.Button("Start AI Diagnosis", variant="primary")
                diagnosis_output = gr.HTML()
                probability_chart = gr.Plot()

            with gr.Tab("Grad-CAM"):
                gradcam_image = gr.Image(label="Explainability Heatmap")

            with gr.Tab("Segmentation"):
                segmentation_panel = gr.Image(label="Tumor Segmentation")

            with gr.Tab("Report"):
                report_button = gr.Button("Generate PDF Report")
                report_file = gr.File(label="Download Report", visible=False)

    return demo
""",
    )

    doc.add_heading("9. Security Features", level=1)
    add_snippet(
        doc,
        "src/api/fastapi_app.py",
        "Input validation and service availability checks.",
        "Prevents invalid image uploads and unavailable model states from crashing the API.",
        """
if _engine is None:
    raise HTTPException(
        status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
        detail="Model not loaded. Check server logs.",
    )

try:
    contents = await file.read()
    image = Image.open(io.BytesIO(contents)).convert("RGB")
    image_np = np.array(image)
except Exception as e:
    raise HTTPException(
        status_code=422,
        detail=f"Cannot decode image: {e}"
    )
""",
    )

    doc.add_heading("10. Testing and Validation", level=1)
    add_snippet(
        doc,
        "src/evaluation/metrics.py",
        "Computes accuracy, AUC, sensitivity, specificity, F1-score, and confusion matrix.",
        "Provides the core academic evaluation metrics for model performance.",
        """
def compute_all_metrics(y_true, y_pred, y_prob, num_classes, class_names):
    accuracy = accuracy_score(y_true, y_pred)
    auc = roc_auc_score(y_true, y_prob, multi_class="ovr", average="macro")
    sensitivity = recall_score(y_true, y_pred, average=None, zero_division=0)
    specificity = compute_specificity(y_true, y_pred, num_classes)
    f1 = f1_score(y_true, y_pred, average="macro", zero_division=0)
    confusion = confusion_matrix(y_true, y_pred, labels=list(range(num_classes)))

    return {
        "accuracy": float(accuracy),
        "auc": float(auc),
        "sensitivity": float(sensitivity.mean()),
        "specificity": float(specificity.mean()),
        "f1_macro": float(f1),
        "confusion_matrix": confusion.tolist(),
    }
""",
    )

    doc.add_heading("11. Deployment", level=1)
    add_snippet(
        doc,
        "Dockerfile",
        "Containerized runtime for API and AI dependencies.",
        "Supports reproducible deployment with PyTorch, CUDA, and project dependencies.",
        """
FROM pytorch/pytorch:2.1.2-cuda11.8-cudnn8-runtime AS base

RUN apt-get update && apt-get install -y --no-install-recommends \\
    libgl1-mesa-glx libglib2.0-0 libsm6 libxext6 libxrender-dev \\
    libgomp1 curl git && rm -rf /var/lib/apt/lists/*

WORKDIR /app

COPY requirements.txt .
RUN pip install --no-cache-dir --upgrade pip && \\
    pip install --no-cache-dir -r requirements.txt

COPY . .

ENV PYTHONPATH=/app
ENV PYTHONUNBUFFERED=1
ENV CONFIG_PATH=/app/config.yaml
ENV CHECKPOINT_DIR=/app/checkpoints

EXPOSE 8000 7860

CMD ["uvicorn", "src.api.fastapi_app:app", "--host", "0.0.0.0", "--port", "8000"]
""",
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
